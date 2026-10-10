import base64
import csv
import io
import json
import random
import sys
import tempfile
import threading
import unittest
import urllib.request
from collections import Counter
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import reconcile as app
import generate_samples as fixtures

def source(rows, header=1, names=None):
    wb=Workbook()
    ws=wb.active
    ws.title='数据'
    for _ in range(header-1): ws.append(['说明'])
    ws.append(names or ['编号','供应商','金额'])
    for row in rows: ws.append(row)
    b=io.BytesIO();wb.save(b)
    return {'data':base64.b64encode(b.getvalue()).decode(),'name':'test.xlsx','sheet':'数据',
            'header':header,'key':0,'supplier':1,'amount':2}

def run(a,b,mode='strict',tolerance='0.01'):
    return app.reconcile({'left':source(a),'right':source(b),'mode':mode,'tolerance':tolerance})

class RuleTests(unittest.TestCase):
    def test_pagination_search_and_missing_key_issues(self):
        rows=[[str(i),'一致','','','S','S',1,1,0] for i in range(125)]
        rows.append(['late-key','金额差异','8','9','late-supplier','S',2,1,1])
        issues=[['A',10,'','S','','编号为空']]
        token=app.cache_result(rows,issues)
        out=app.result_page({'token':token,'filter':'all','page':3})
        self.assertEqual((out['count'],out['pages'],len(out['rows'])),(126,3,26))
        found=app.result_page({'token':token,'filter':'all','query':'late-supplier'})
        self.assertEqual(found['rows'][0][0],'late-key')
        self.assertEqual(app.result_page({'token':token,'filter':'issues','query':'编号为空'})['rows'][0][2],'')
        self.assertEqual(app.result_page({'token':token,'filter':'all','query':'absent'})['count'],0)
        with self.assertRaises(ValueError): app.result_page({'token':token,'filter':'bad'})
        with self.assertRaises(ValueError): app.result_page({'token':token,'page':0})

    def test_result_cache_eviction_and_expiry(self):
        old=app.cache_result([],[])
        app.cache_result([],[]);app.cache_result([],[])
        with self.assertRaises(ValueError): app.result_page({'token':old})
        token=app.cache_result([],[])
        with app.RESULT_LOCK: app.RESULT_CACHE[token]=(app.time.monotonic()-1801,[],[])
        with self.assertRaises(ValueError): app.result_page({'token':token})

    def test_source_record_trace_and_limit(self):
        a=source([['001','甲',10],['002','乙','=1+1'],['003','丙',30]],header=3)
        out=app.source_records(a,[4,6])
        self.assertEqual(out['headers'],['编号','供应商','金额'])
        self.assertEqual(out['rows'],[{'row':4,'values':['001','甲','10']},{'row':6,'values':['003','丙','30']}])
        with self.assertRaises(ValueError): app.source_records(a,[3])
        with self.assertRaises(ValueError): app.source_records(a,list(range(4,55)))
        self.assertEqual(app.source_records(a,[])['rows'],[])

    def test_csv_source_record_trace(self):
        a={'data':base64.b64encode('说明\n编号,供应商,金额\n001,甲,12\n002,乙,20\n'.encode('gb18030')).decode(),'name':'a.csv','header':2}
        self.assertEqual(app.source_records(a,[4])['rows'],[{'row':4,'values':['002','乙','20']}])

    def test_filters_count_full_result_beyond_preview(self):
        rows=[[str(i),'一致','','','S','S',1,1,0] for i in range(250)]
        rows.append(['last','金额差异；重复编号待确认','2','3','S','S',2,1,1])
        buckets=app.result_buckets(rows)
        self.assertEqual(buckets['all']['count'],251)
        self.assertEqual(len(buckets['all']['rows']),200)
        self.assertEqual(buckets['amount']['rows'][0][0],'last')
        self.assertEqual(buckets['duplicate']['count'],1)
        self.assertEqual(buckets['review']['count'],1)

    def test_source_preview_rows_and_formula_validation(self):
        a=source([['001','S','1,200.00'],['002','S','=1+1'],['003','S','bad'],['004','S',None],['005','S',0],['006','S',6]],header=3)
        out=app.inspect(a['data'],a['name'],3,with_preview=True)
        rows=out['previews']['数据']
        self.assertEqual(len(rows),5)
        self.assertEqual([r['row'] for r in rows],[4,5,6,7,8])
        self.assertEqual(rows[0]['values'][0],'001')
        self.assertEqual([r['amountValid'][2] for r in rows],[True,False,False,False,True])
        self.assertEqual(rows[1]['formulas'],[2])

    def test_csv_source_preview(self):
        raw='说明\n编号,供应商,金额\n001,甲,12.34\n002,乙,错误\n'.encode('gb18030')
        out=app.inspect(base64.b64encode(raw).decode(),'sample.csv',2,with_preview=True)
        self.assertEqual(out['sheets']['CSV'],['编号','供应商','金额'])
        rows=out['previews']['CSV']
        self.assertEqual(rows[0]['row'],3)
        self.assertEqual(rows[0]['values'],['001','甲','12.34'])
        self.assertFalse(rows[1]['amountValid'][2])

    def test_matching(self):
        _,rows,issues=run([['A','S',1]],[['A','S','1.00']])
        self.assertEqual(rows[0][1],'一致');self.assertFalse(issues)

    def test_missing_and_different(self):
        _,rows,_=run([['A','S',1],['B','S',2]],[['A','S',3],['C','S',4]])
        self.assertEqual([r[1] for r in rows],['金额差异','仅A存在','仅B存在'])

    def test_duplicate_modes(self):
        a=[['A','S',40],['A','S',60]];b=[['A','S',100]]
        self.assertEqual(run(a,b)[1][0][1],'重复编号待确认')
        self.assertEqual(run(a,b,'sum')[1][0][1],'一致')
        self.assertEqual(run([['A','S',100]]*2,b,'sum')[1][0][1],'金额差异')

    def test_supplier_and_amount_both_visible(self):
        status=run([['A','S',1]],[['A','T',2]])[1][0][1]
        self.assertEqual(status,'供应商不一致；金额差异')

    def test_missing_even_when_duplicate(self):
        status=run([['A','S',1]]*2,[])[1][0][1]
        self.assertEqual(status,'重复编号待确认；仅A存在')

    def test_multisupplier_not_matched(self):
        status=run([['A','S',1],['A','T',2]],[['A','S',3]],'sum')[1][0][1]
        self.assertIn('同编号包含多个供应商',status)

    def test_invalid_blank_and_zero(self):
        _,rows,issues=run([['A','S',None],['B','S','oops'],[None,'S',5]],[['A','S',0],['B','S',0]])
        self.assertEqual(len(issues),3)
        self.assertTrue(all('数据异常' in r[1] for r in rows))

    def test_blank_supplier(self):
        self.assertIn('数据异常',run([['A',None,1]],[['A','S',1]])[1][0][1])

    def test_tolerance(self):
        self.assertEqual(run([['A','S','0.30']],[['A','S','0.29']])[1][0][1],'一致')
        self.assertEqual(run([['A','S','0.30']],[['A','S','0.289']])[1][0][1],'金额差异')

    def test_refunds_thousands(self):
        self.assertEqual(run([['A','S','(1,234.56)']],[['A','S',-1234.56]])[1][0][1],'一致')
        with self.assertRaises(ValueError): app.amount('1,2')
        with self.assertRaises(ValueError): app.amount('NaN')
        with self.assertRaises(ValueError): app.amount(True)

    def test_numeric_precision_limits(self):
        with self.assertRaises(ValueError): app.amount('1e1000')
        with self.assertRaises(ValueError): app.amount('0.0000001')

    def test_identifiers(self):
        _,rows,_=run([['001','S',1],['a','S',1]],[['1','S',1],['A','S',1]])
        self.assertEqual(len(rows),4)
        self.assertFalse(any(r[1]=='一致' for r in rows))

    def test_trim(self):
        self.assertEqual(run([[' A ',' S ',1]],[['A','S',1]])[1][0][1],'一致')

    def test_formula_detection_and_text_safety(self):
        raw,rows,issues=run([['A','S','=1+1']],[['A','S',2]])
        self.assertIn('数据异常',rows[0][1]);self.assertTrue(issues)
        wb=load_workbook(io.BytesIO(raw))
        self.assertEqual(wb['数据异常']['E2'].data_type,'s')

    def test_empty(self):
        _,rows,issues=run([],[])
        self.assertEqual(rows,[]);self.assertEqual(issues,[])

    def test_negative_tolerance_and_bad_mapping(self):
        with self.assertRaises(ValueError): run([],[],tolerance='-1')
        a=source([]);a['supplier']=0
        with self.assertRaises(ValueError): app.reconcile({'left':a,'right':source([])})

    def test_header_offset(self):
        a=source([['A','S',1]],header=3)
        self.assertEqual(app.inspect(a['data'],header=3)['数据'],['编号','供应商','金额'])
        rows=app.reconcile({'left':a,'right':a})[1]
        self.assertEqual(rows[0][2],'4');self.assertEqual(rows[0][1],'一致')

    def test_composite_keys_no_collision(self):
        a=source([['A',1,'S',10],['A',2,'S',20]],names=['订单','行号','供应商','金额'])
        a.update(keys=[0,1],supplier=2,amount=3)
        rows=app.reconcile({'left':a,'right':a})[1]
        self.assertEqual(len(rows),2);self.assertTrue(all(r[1]=='一致' for r in rows))

    def test_csv_encodings(self):
        for encoding in ('utf-8-sig','gb18030'):
            b=io.StringIO();w=csv.writer(b);w.writerows([['编号','供应商','金额'],['001','供应商甲','1,000.00']])
            a={'data':base64.b64encode(b.getvalue().encode(encoding)).decode(),'name':'sample.csv','sheet':'CSV','key':0,'supplier':1,'amount':2}
            self.assertEqual(app.inspect(a['data'],a['name'])['CSV'],['编号','供应商','金额'])
            self.assertEqual(app.reconcile({'left':a,'right':a})[1][0][1],'一致')

    def test_invalid_workbook(self):
        with self.assertRaises(ValueError): app.book(base64.b64encode(b'not excel').decode())

    def test_export_sheet_trace_and_string_safety(self):
        raw,_,_=run([['A','+SUM','1.25']],[['A','+SUM','2.25']])
        wb=load_workbook(io.BytesIO(raw),data_only=False)
        self.assertEqual(wb.sheetnames,['核对摘要','差异清单','全部核对','重复编号','数据异常'])
        self.assertEqual(wb['差异清单']['C2'].value,'2')
        self.assertEqual(wb['差异清单']['E2'].data_type,'s')
        self.assertEqual(wb['差异清单']['I2'].value,-1)

    def test_decimal_results_unchanged_by_export(self):
        _,rows,_=run([['A','S','0.30']],[['A','S','0.29']])
        self.assertEqual(rows[0][-1],Decimal('0.01'))

    def test_row_order_invariance(self):
        rng=random.Random(18)
        for _ in range(10):
            a,b,expected,issues=fixtures.build(80,'mixed',rng.randrange(10000))
            before=run(a,b)[1]
            rng.shuffle(a);rng.shuffle(b)
            after=run(a,b)[1]
            self.assertEqual([(r[0],r[1],r[-1]) for r in before],[(r[0],r[1],r[-1]) for r in after])

class SampleTests(unittest.TestCase):
    def test_all_samples_against_independent_answers(self):
        samples=ROOT/'samples'
        if not (samples/'index.json').exists(): fixtures.generate(samples)
        index=json.loads((samples/'index.json').read_text(encoding='utf-8'))
        self.assertEqual(len(index),12)
        for sample in index:
            answer=json.loads((samples/sample['answer']).read_text(encoding='utf-8'))
            configs={}
            for side in ('left','right'):
                path=samples/sample[side]
                configs[side]={'data':base64.b64encode(path.read_bytes()).decode(),'name':path.name,'sheet':'数据','key':0,'supplier':1,'amount':2}
            for mode in ('strict','sum'):
                with self.subTest(sample=sample['id'],mode=mode):
                    _,rows,issues=app.reconcile({**configs,'mode':mode,'tolerance':'0.01'})
                    self.assertEqual(len(issues),answer['issues'])
                    self.assertEqual(dict(Counter(r[1] for r in rows)),answer['status_counts'][mode])
                    self.assertEqual({r[0] for r in rows},set(answer['expected']))
                    for r in rows:
                        e=answer['expected'][r[0]]
                        self.assertEqual(r[1],e[mode])
                        for actual,name in [(r[6],'a'),(r[7],'b'),(r[8],'difference')]:
                            self.assertEqual(actual,Decimal(e[name]) if e[name] is not None else None)

class ServerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server=app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
        cls.url=f'http://127.0.0.1:{cls.server.server_port}'
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown();cls.server.server_close();cls.thread.join()

    def post(self,path,value,origin=None):
        headers={'Content-Type':'application/json','Origin':origin or self.url}
        request=urllib.request.Request(self.url+path,json.dumps(value).encode(),headers)
        return urllib.request.urlopen(request)

    def test_page_and_samples(self):
        with urllib.request.urlopen(self.url) as r: self.assertIn('id="cards"',r.read().decode())
        with urllib.request.urlopen(self.url+'/samples') as r: self.assertEqual(len(json.load(r)),12)

    def test_inspect_and_compare_http(self):
        a=source([['A','S',1]])
        with self.post('/inspect',{'data':a['data'],'name':a['name']}) as r:
            self.assertIn('数据',json.load(r)['sheets'])
        with self.post('/compare',{'left':a,'right':a}) as r:
            out=json.load(r);self.assertEqual(out['matched'],1)
            self.assertEqual(out['buckets']['all']['count'],1)
            self.assertEqual(load_workbook(io.BytesIO(base64.b64decode(out['file']))).sheetnames[0],'核对摘要')

    def test_source_http(self):
        a=source([['001','S',12]])
        with self.post('/source',{'source':a,'rows':[2]}) as r:
            self.assertEqual(json.load(r)['rows'][0]['values'],['001','S','12'])

    def test_foreign_origin_rejected(self):
        with self.assertRaises(urllib.error.HTTPError) as e:
            self.post('/inspect',{},origin='https://example.com')
        self.assertEqual(e.exception.code,400)

    def test_health_and_deployment_origin(self):
        with urllib.request.urlopen(self.url+'/health') as r:
            self.assertEqual(json.load(r), {'status':'ok'})
        previous = app.Handler.allowed_origins
        app.Handler.allowed_origins = {'https://reconcile.internal'}
        try:
            a = source([['A','S',1]])
            with self.post('/compare', {'left':a,'right':a}, origin='https://reconcile.internal') as r:
                self.assertEqual(json.load(r)['matched'],1)
            with self.assertRaises(urllib.error.HTTPError):
                self.post('/inspect', {}, origin=self.url)
        finally:
            app.Handler.allowed_origins = previous

    def test_sample_traversal_rejected(self):
        with self.assertRaises(urllib.error.HTTPError) as e:
            urllib.request.urlopen(self.url+'/samples/../requirements.txt')
        self.assertEqual(e.exception.code,404)

if __name__=='__main__': unittest.main()
