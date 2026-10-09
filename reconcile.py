#!/usr/bin/env python3
"""本机 Excel 核对工具。仅依赖 openpyxl；不修改输入文件。"""
import base64
import argparse
import csv
import io
import json
import re
import threading
import webbrowser
import zipfile
from urllib.parse import unquote
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell import WriteOnlyCell

ROOT = Path(__file__).resolve().parent
VERSION = '1.0.0'
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100000

def decoded(encoded):
    try:
        raw = base64.b64decode(encoded, validate=True)
    except Exception:
        raise ValueError('文件编码无效，请重新选择文件')
    if len(raw) > MAX_FILE_BYTES:
        raise ValueError('单个文件不能超过20MB')
    return raw

def book(encoded):
    raw = decoded(encoded)
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(v.file_size for v in z.infolist()) > 200 * 1024 * 1024:
                raise ValueError('Excel解压后超过200MB，请拆分文件')
        return load_workbook(io.BytesIO(raw), data_only=False, read_only=True)
    except ValueError:
        raise
    except Exception:
        raise ValueError('无法读取.xlsx文件，请检查格式、损坏或加密情况')

def csv_rows(config):
    raw = decoded(config['data'])
    for encoding in ('utf-8-sig', 'gb18030'):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            pass
    else:
        raise ValueError('CSV须使用UTF-8或GB18030编码')
    return csv.reader(io.StringIO(text))

def header_number(config):
    n = int(config.get('header', 1))
    if not 1 <= n <= 100:
        raise ValueError('表头行须在1到100之间')
    return n

def inspect(encoded, name='file.xlsx', header=1):
    config = {'data': encoded, 'header': header}
    n = header_number(config)
    if name.lower().endswith('.csv'):
        rows = csv_rows(config)
        headers = next((row for i, row in enumerate(rows, 1) if i == n), [])
        if not headers:
            raise ValueError('指定表头行为空')
        return {'CSV': [str(v).strip() for v in headers]}
    if not name.lower().endswith('.xlsx'):
        raise ValueError('只支持.xlsx或.csv文件')
    wb = book(encoded)
    result = {}
    try:
        for ws in wb:
            headers = next(ws.iter_rows(min_row=n, max_row=n, values_only=True), ())
            result[ws.title] = [str(v).strip() if v is not None else '' for v in headers]
    finally:
        wb.close()
    return result

def amount(value):
    if value is None or str(value).strip() == '':
        raise ValueError('金额为空')
    if isinstance(value, bool):
        raise ValueError('金额不能是布尔值')
    s = str(value).strip().replace('，', ',')
    if s.startswith('(') and s.endswith(')'):
        s = '-' + s[1:-1]
    if ',' in s and not re.fullmatch(r'[+-]?\d{1,3}(,\d{3})+(\.\d+)?', s):
        raise ValueError('千分位格式无效')
    if not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?', s.replace(',', '')):
        raise ValueError('金额不是有效数字')
    s = s.replace(',', '')
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise ValueError('金额不是有效数字')
    if not d.is_finite():
        raise ValueError('金额不是有限数字')
    if abs(d) >= Decimal('1e12') or d.as_tuple().exponent < -6:
        raise ValueError('金额须小于1万亿元且最多6位小数')
    return d

def read_source(config, label):
    n = header_number(config)
    csv_mode = config.get('name', '').lower().endswith('.csv')
    wb = None if csv_mode else book(config['data'])
    groups = defaultdict(list)
    issues = []
    total = Decimal(0)
    count = 0
    try:
        if csv_mode:
            source = csv_rows(config)
            headers = next((row for i, row in enumerate(source, 1) if i == n), [])
            source = ((row, set()) for row in source)
        else:
            try:
                ws = wb[config['sheet']]
            except KeyError:
                raise ValueError(f'{label}工作表不存在，请重新选择')
            headers = next(ws.iter_rows(min_row=n, max_row=n, values_only=True), ())
            source = ((tuple(c.value for c in row), {i for i,c in enumerate(row) if c.data_type == 'f'})
                      for row in ws.iter_rows(min_row=n+1))
        key_indices = [int(v) for v in config.get('keys', [config.get('key', 0)])]
        indices = key_indices + [int(config['supplier']), int(config['amount'])]
        width = len(headers)
        if not key_indices or any(i < 0 or i >= width for i in indices) or len(set(indices)) != len(indices):
            raise ValueError('编号、供应商、金额须选择不同且有效的列')
        for rowno, (cells, formula_indices) in enumerate(source, n+1):
            if rowno - n > MAX_ROWS:
                raise ValueError('单表最多处理10万行；请清理多余格式行或拆分文件')
            if all(v is None or str(v).strip() == '' for v in cells):
                continue
            count += 1
            raw_keys = [cells[i] if i < len(cells) else None for i in key_indices]
            key_parts = tuple(str(k).strip() if k is not None else '' for k in raw_keys)
            key = key_parts[0] if len(key_parts) == 1 else json.dumps(key_parts, ensure_ascii=False)
            supplier, money = [cells[i] if i < len(cells) else None for i in indices[-2:]]
            sup = str(supplier).strip() if supplier is not None else ''
            errors = []
            if not all(key_parts):
                errors.append('编号为空')
            if any(isinstance(k, bool) or not isinstance(k, (str, int, float, type(None))) for k in raw_keys):
                errors.append('编号须为文本或数字')
            if not sup:
                errors.append('供应商为空')
            try:
                d = amount(money)
                total += d
            except ValueError as e:
                d = None
                errors.append(str(e))
            if formula_indices.intersection(indices):
                errors.append('核对字段含公式，请在副本中转为值')
                if d is not None:
                    total -= d
                d = None
            if errors:
                issues.append([label, rowno, key, sup, str(money) if money is not None else '', '；'.join(errors)])
            # 保留问题编号，阻止把另一侧误判成仅单侧存在。
            if all(key_parts):
                groups[key].append({'row': rowno, 'supplier': sup, 'amount': d, 'invalid': bool(errors)})
    finally:
        if wb:
            wb.close()
    return groups, issues, total, count

def reconcile(payload):
    tolerance = amount(payload.get('tolerance', '0.01'))
    if tolerance < 0:
        raise ValueError('金额容差不能为负数')
    mode = payload.get('mode', 'strict')
    if mode not in ('strict', 'sum'):
        raise ValueError('未知核对模式')
    left, li, lt, lc = read_source(payload['left'], 'A')
    right, ri, rt, rc = read_source(payload['right'], 'B')
    results, duplicates = [], []
    for side, groups in [('A', left), ('B', right)]:
        for key, rows in groups.items():
            if len(rows) > 1:
                for r in rows:
                    duplicates.append([side, key, r['row'], r['supplier'], r['amount']])
    for key in sorted(set(left) | set(right)):
        a, b = left.get(key, []), right.get(key, [])
        reasons = []
        invalid = any(r['invalid'] for r in a + b)
        duplicate = mode == 'strict' and (len(a) > 1 or len(b) > 1)
        mixed = any(len({r['supplier'] for r in rows}) > 1 for rows in (a, b))
        if invalid:
            reasons.append('数据异常待处理')
        if duplicate:
            reasons.append('重复编号待确认')
        if mixed:
            reasons.append('同编号包含多个供应商')
        av = sum((r['amount'] for r in a), Decimal(0)) if a and not any(r['invalid'] for r in a) else None
        bv = sum((r['amount'] for r in b), Decimal(0)) if b and not any(r['invalid'] for r in b) else None
        diff = av - bv if av is not None and bv is not None else None
        if not b:
            reasons.append('仅A存在')
        elif not a:
            reasons.append('仅B存在')
        else:
            if all(r['supplier'] for r in a+b) and {r['supplier'] for r in a} != {r['supplier'] for r in b}:
                reasons.append('供应商不一致')
            if diff is not None and abs(diff) > tolerance:
                reasons.append('金额差异')
        status = '；'.join(reasons) if reasons else '一致'
        results.append([key, status, ', '.join(str(r['row']) for r in a), ', '.join(str(r['row']) for r in b),
                        ' / '.join(sorted({r['supplier'] for r in a})), ' / '.join(sorted({r['supplier'] for r in b})), av, bv, diff])
    summary = [
        ['项目', '值'], ['说明', '个人演示工具；不修改原文件；金额单位、币种、含税口径须自行确认一致'],
        ['模式', '重复编号待确认' if mode == 'strict' else '按编号汇总金额（不自动删除重复行）'],
        ['金额容差', tolerance], ['A非空行数', lc], ['B非空行数', rc],
        ['A有效数字金额合计（含编号缺失行）', lt], ['B有效数字金额合计（含编号缺失行）', rt],
        ['有效数字金额合计差（A-B）', lt-rt], ['两侧编号并集数', len(results)],
        ['一致编号数', sum(r[1] == '一致' for r in results)],
        ['待核查编号数', sum(r[1] != '一致' for r in results)], ['异常源数据行数', len(li+ri)],
        ['结论', '存在待核查项' if li+ri or any(r[1] != '一致' for r in results) else '当前规则下核对一致'],
        ['编号规则', '区分大小写；仅去首尾空格；多列编号按组合匹配；文本001与数字1不合并'],
        ['金额规则', '支持逗号千分位、负数及括号负数；空金额不当作0；公式不计算，需转为值'],
        ['差额规则', 'A金额-B金额；绝对差大于容差才标记金额差异；供应商单独比较'],
        ['异常规则', '含无效字段或多供应商的编号不判为一致；无编号行见数据异常'],
        ['文件A', payload['left'].get('name', '')], ['文件B', payload['right'].get('name', '')],
        ['版本', VERSION], ['A表头行', header_number(payload['left'])], ['B表头行', header_number(payload['right'])]
    ]
    wb = Workbook(write_only=True)
    tables = [('核对摘要', summary),
              ('差异清单', [['编号','状态','A源行','B源行','A供应商','B供应商','A金额','B金额','差額A-B']] + [r for r in results if r[1] != '一致']),
              ('全部核对', [['编号','状态','A源行','B源行','A供应商','B供应商','A金额','B金额','差額A-B']] + results),
              ('重复编号', [['来源','编号','源行','供应商','金额']] + duplicates),
              ('数据异常', [['来源','源行','编号','供应商','原金额','原因']] + li + ri)]
    for name, rows in tables:
        ws = wb.create_sheet(name)
        ws.sheet_format.defaultRowHeight = 60 if name == '核对摘要' else 45
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{chr(64+len(rows[0]))}{len(rows)}'
        for i in range(len(rows[0])):
            ws.column_dimensions[chr(65+i)].width = (65 if i == 1 else 35) if name == '核对摘要' else (30 if i == 1 else 22)
        for rowno, row in enumerate(rows, 1):
            output = []
            for i, v in enumerate(row):
                c = WriteOnlyCell(ws, value=float(v) if isinstance(v, Decimal) else v)
                # 导出外部文本为普通字符串，避免公式注入。
                if isinstance(c.value, str):
                    c.data_type = 's'
                c.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(v, Decimal):
                    c.number_format = '#,##0.00####;[Red]-#,##0.00####'
                if rowno == 1:
                    c.fill = PatternFill('solid', fgColor='16324F')
                    c.font = Font(color='FFFFFF', bold=True)
                output.append(c)
            ws.append(output)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), results, li+ri

class Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def do_GET(self):
        if self.path == '/':
            raw = (ROOT / 'interface.html').read_bytes()
            content_type = 'text/html; charset=utf-8'
        elif self.path == '/samples':
            index = ROOT / 'samples' / 'index.json'
            raw = index.read_bytes() if index.exists() else b'[]'
            content_type = 'application/json; charset=utf-8'
        elif self.path.startswith('/samples/'):
            relative = unquote(self.path.removeprefix('/samples/'))
            target = (ROOT / 'samples' / relative).resolve()
            if not target.is_relative_to((ROOT / 'samples').resolve()) or not target.is_file() or target.suffix not in ('.xlsx','.csv','.json'):
                self.send_error(404)
                return
            raw = target.read_bytes()
            content_type = 'application/octet-stream'
        else:
            self.send_error(404)
            return
        self.send_response(200)
        self.send_header('Content-Type', content_type)
        self.send_header('Cache-Control', 'no-store')
        self.end_headers()
        self.wfile.write(raw)

    def do_POST(self):
        try:
            # 只接受本机页面发起的请求。
            origin = self.headers.get('Origin', '')
            expected = 'http://' + self.headers.get('Host', '')
            if origin != expected or not expected.startswith('http://127.0.0.1:'):
                raise ValueError('请求来源不允许')
            size = int(self.headers.get('Content-Length', 0))
            if size <= 0 or size > 60 * 1024 * 1024:
                raise ValueError('请求大小超出限制')
            payload = json.loads(self.rfile.read(size))
            if self.path == '/inspect':
                result = {'sheets': inspect(payload['data'], payload.get('name','file.xlsx'), payload.get('header',1))}
            elif self.path == '/compare':
                raw, rows, issues = reconcile(payload)
                result = {'file': base64.b64encode(raw).decode(), 'matched': sum(r[1]=='一致' for r in rows),
                          'review': sum(r[1]!='一致' for r in rows), 'issues': len(issues),
                          'preview': [[str(v) if v is not None else '' for v in r] for r in rows[:20]]}
            else:
                raise ValueError('未知操作')
            data = json.dumps(result, ensure_ascii=False).encode()
            self.send_response(200)
        except Exception as e:
            data = json.dumps({'error': str(e)}, ensure_ascii=False).encode()
            self.send_response(400)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.end_headers()
        self.wfile.write(data)

def main():
    parser = argparse.ArgumentParser(description='本机Excel/CSV两表核对工具')
    parser.add_argument('--port', type=int, default=0, help='端口，默认自动选择空闲端口')
    parser.add_argument('--no-browser', action='store_true', help='启动时不自动打开浏览器')
    parser.add_argument('--config', type=Path, help='以JSON配置进行批量核对，不启动页面')
    parser.add_argument('--output', type=Path, default=Path('核对结果.xlsx'))
    args = parser.parse_args()
    if not (ROOT / 'samples' / 'index.json').exists():
        print('首次运行，正在生成模拟练习样本和标准答案…', flush=True)
        from generate_samples import generate
        generate(ROOT / 'samples')
    if args.config:
        payload = json.loads(args.config.read_text(encoding='utf-8'))
        for side in ('left','right'):
            source = payload[side]
            path = (args.config.parent / source.pop('path')).resolve()
            if args.output.resolve() == path:
                parser.error('输出文件不能覆盖任何输入文件')
            source['name'] = path.name
            source['data'] = base64.b64encode(path.read_bytes()).decode()
        raw, rows, issues = reconcile(payload)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_bytes(raw)
        print(f'已导出 {args.output}：一致{sum(r[1]=="一致" for r in rows)}个编号，异常源行{len(issues)}行')
        return
    server = ThreadingHTTPServer(('127.0.0.1', args.port), Handler)
    url = f'http://127.0.0.1:{server.server_port}'
    print('工具已启动：' + url, flush=True)
    print('关闭工具：回到此终端按 Control+C。', flush=True)
    if not args.no_browser:
        threading.Timer(0.6, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        server.server_close()

if __name__ == '__main__':
    main()
