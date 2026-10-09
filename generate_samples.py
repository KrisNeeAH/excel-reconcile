"""生成有独立标准答案的模拟练习样本，不调用核对引擎。"""
import argparse
import json
import random
from collections import Counter
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent
SEED = 20261009
SCENARIOS = [
    ('01_intro', '入门：常见差异', 10, 'intro', '十个业务编号，涵盖缺失、差额、多行、重复、无效金额与退款。'),
    ('02_clean', '正常订单', 2000, 'clean', '两表顺序不同，编号与金额应全部一致。'),
    ('03_amount', '金额差异', 2000, 'amount', '每十个编号中一个有明确金额差异。'),
    ('04_missing', '单侧缺失', 2000, 'missing', 'A和B分别缺失不同记录，核对单侧存在项。'),
    ('05_duplicates', '重复录入', 1000, 'duplicates', '部分A记录被完整重复录入；汇总不会修正重复错误。'),
    ('06_split', '多行订单', 1000, 'split', '部分A订单拆成两行；确认业务口径后可按编号汇总。'),
    ('07_supplier', '供应商冲突', 1000, 'supplier', '部分两侧供应商不一致，或同一编号有多个供应商。'),
    ('08_dirty', '脏数据', 1000, 'dirty', '覆盖空编号、空金额、错误金额和缺失供应商。'),
    ('09_tolerance', '容差边界', 500, 'tolerance', '交替出现0.01和0.02差额，默认容差0.01。'),
    ('10_refunds', '退款与千分位', 1000, 'refunds', '负数、括号负数、千分位文本金额的核对。'),
    ('11_identifiers', '前导零与大小写', 300, 'identifiers', '001与1、大小写不同编号不应自动合并。'),
    ('12_large', '大规模混合场景', 10000, 'mixed', '一万个业务编号，混合正常、金额差异、缺失、多行、重复和异常。'),
]
TYPES = ['normal','difference','only_a','only_b','split','duplicate','supplier','invalid','empty','refund']

def write_xlsx(path, rows):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet('数据')
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 20
    ws.freeze_panes = 'A2'
    from openpyxl.cell import WriteOnlyCell
    heads = []
    for v in ['订单编号','供应商','金额']:
        c = WriteOnlyCell(ws, value=v)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='16324F')
        heads.append(c)
    ws.append(heads)
    for row in rows:
        out = []
        for v in row:
            c = WriteOnlyCell(ws, value=float(v) if isinstance(v, Decimal) else v)
            if isinstance(v,str):
                c.data_type = 's'
            if isinstance(v,Decimal):
                c.number_format = '#,##0.00'
            out.append(c)
        ws.append(out)
    wb.save(path)

def kind_for(name, i):
    if name == 'intro': return TYPES[i % len(TYPES)]
    if name == 'clean': return 'normal'
    if name == 'amount': return 'difference' if i % 10 == 0 else 'normal'
    if name == 'missing': return ['only_a','only_b','normal','normal','normal'][i % 5]
    if name == 'duplicates': return 'duplicate' if i % 5 == 0 else 'normal'
    if name == 'split': return 'split' if i % 2 == 0 else 'normal'
    if name == 'supplier': return ['supplier','mixed_supplier'][i % 2] if i % 5 == 0 else 'normal'
    if name == 'dirty': return ['invalid','empty','blank_key','blank_supplier'][i % 4] if i % 5 == 0 else 'normal'
    if name == 'tolerance': return 'edge' if i % 2 == 0 else 'outside'
    if name == 'refunds': return 'refund' if i % 2 == 0 else 'thousands'
    if name == 'identifiers': return ['leading_zero','case','normal'][i % 3]
    return (['normal']*7 + TYPES[1:] + ['mixed_supplier','blank_key','blank_supplier'])[i % 19]

def build(count, scenario, seed):
    """答案由注入场景预先声明，不以核对引擎的输出作答案。"""
    rng = random.Random(seed)
    a, b, expected = [], [], {}
    issue_count = 0
    for i in range(count):
        key = f'PO{i+1:07d}'
        supplier = f'模拟供应商{i % 37 + 1:03d}'
        value = Decimal(rng.randint(10000, 5000000)) / 100
        kind = kind_for(scenario, i)
        strict = summed = '一致'
        av = bv = value
        ar = [[key,supplier,value]]
        br = [[key,supplier,value]]
        if kind == 'difference':
            bv = value - Decimal('12.34')
            br[0][2] = bv
            strict = summed = '金额差异'
        elif kind == 'only_a':
            br=[]; bv=None; strict=summed='仅A存在'
        elif kind == 'only_b':
            ar=[]; av=None; strict=summed='仅B存在'
        elif kind == 'split':
            ar=[[key,supplier,Decimal('40')],[key,supplier,value-40]]
            strict='重复编号待确认'
        elif kind == 'duplicate':
            ar.append(ar[0].copy()); av=value*2
            strict='重复编号待确认；金额差异'; summed='金额差异'
        elif kind == 'supplier':
            br[0][1]='模拟供应商不同'; strict=summed='供应商不一致'
        elif kind == 'mixed_supplier':
            ar=[[key,supplier,Decimal('40')],[key,'模拟供应商不同',value-40]]
            strict='重复编号待确认；同编号包含多个供应商；供应商不一致'
            summed='同编号包含多个供应商；供应商不一致'
        elif kind in ('invalid','empty','blank_supplier'):
            if kind=='invalid': ar[0][2]='错误金额'
            if kind=='empty': ar[0][2]=None
            if kind=='blank_supplier': ar[0][1]=None
            av=None; issue_count+=1; strict=summed='数据异常待处理'
        elif kind == 'blank_key':
            ar[0][0]=None; av=None; issue_count+=1; strict=summed='仅B存在'
        elif kind == 'edge':
            bv=value-Decimal('0.01');br[0][2]=bv
        elif kind == 'outside':
            bv=value-Decimal('0.02');br[0][2]=bv;strict=summed='金额差异'
        elif kind == 'refund':
            av=bv=-value; ar[0][2]=-value;br[0][2]=f'({value:,.2f})'
        elif kind == 'thousands':
            ar[0][2]=f'{value:,.2f}'
        elif kind in ('leading_zero','case'):
            ka=f'00{i+1:06d}' if kind=='leading_zero' else key.lower()
            kb=str(i+1) if kind=='leading_zero' else key
            ar[0][0]=ka;br[0][0]=kb
            expected[ka]={'strict':'仅A存在','sum':'仅A存在','a':str(value),'b':None,'difference':None,'scenario':kind}
            expected[kb]={'strict':'仅B存在','sum':'仅B存在','a':None,'b':str(value),'difference':None,'scenario':kind}
            a.extend(ar);b.extend(br);continue
        expected[key]={'strict':strict,'sum':summed,'a':str(av) if av is not None else None,
                       'b':str(bv) if bv is not None else None,
                       'difference':str(av-bv) if av is not None and bv is not None else None,'scenario':kind}
        a.extend(ar);b.extend(br)
    rng.shuffle(a);rng.shuffle(b)
    return a,b,expected,issue_count

def generate(output=ROOT/'samples'):
    output.mkdir(parents=True,exist_ok=True)
    index=[]
    for j,(slug,title,count,scenario,description) in enumerate(SCENARIOS):
        folder=output/slug
        folder.mkdir(exist_ok=True)
        a,b,expected,issues=build(count,scenario,SEED+j)
        write_xlsx(folder/'A.xlsx',a)
        write_xlsx(folder/'B.xlsx',b)
        counts={mode:dict(Counter(item[mode] for item in expected.values())) for mode in ('strict','sum')}
        answer={'source':'人工定义场景＋确定性模拟数据；非客户数据；不调用核对引擎生成答案',
                'seed':SEED+j,'records':count,'rows_A':len(a),'rows_B':len(b),'tolerance':'0.01',
                'issues':issues,'status_counts':counts,'expected':expected}
        (folder/'answer.json').write_text(json.dumps(answer,ensure_ascii=False,indent=2),encoding='utf-8')
        answer_book = Workbook()
        summary = answer_book.active
        summary.title = '标准摘要'
        summary.append(['项目','标准答案'])
        for row in [('样本',title),('数据来源','固定种子模拟数据；不包含真实客户数据'),('业务编号场景数',count),
                    ('A源行数',len(a)),('B源行数',len(b)),('异常源行数',issues),('金额容差',0.01)]:
            summary.append(row)
        for mode,label in [('strict','默认模式'),('sum','汇总模式')]:
            for status,number in counts[mode].items():
                summary.append([label+' / '+status,number])
        details = answer_book.create_sheet('逐项答案')
        details.append(['编号','默认模式状态','汇总模式状态','A有效核对金额','B有效核对金额','差额A-B','注入场景'])
        for key,e in sorted(expected.items()):
            details.append([key,e['strict'],e['sum'],*[float(e[k]) if e[k] is not None else None for k in ('a','b','difference')],e['scenario']])
        for ws in answer_book:
            ws.sheet_format.defaultRowHeight=36
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for row in ws:
                for c in row:
                    c.alignment=Alignment(vertical='top',wrap_text=True)
            for c in ws[1]:
                c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='16324F')
            for i in range(1,ws.max_column+1):
                ws.column_dimensions[ws.cell(1,i).column_letter].width=48 if i in (2,3) else 24
            if ws==details:
                for row in ws.iter_rows(min_row=2,min_col=4,max_col=6):
                    for c in row: c.number_format='#,##0.00'
        answer_book.save(folder/'标准答案.xlsx')
        config={'left':{'path':'A.xlsx','sheet':'数据','header':1,'key':0,'supplier':1,'amount':2},
                'right':{'path':'B.xlsx','sheet':'数据','header':1,'key':0,'supplier':1,'amount':2},
                'mode':'strict','tolerance':'0.01'}
        (folder/'config.json').write_text(json.dumps(config,ensure_ascii=False,indent=2),encoding='utf-8')
        index.append({'id':slug,'title':title,'records':count,'description':description,
                      'left':slug+'/A.xlsx','right':slug+'/B.xlsx','answer':slug+'/answer.json',
                      'answerWorkbook':slug+'/标准答案.xlsx'})
    (output/'index.json').write_text(json.dumps(index,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'已生成{len(index)}组样本，合计{sum(s["records"] for s in index):,}个业务编号场景')
    return index

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--output',type=Path,default=ROOT/'samples')
    generate(p.parse_args().output)
